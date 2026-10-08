import io
import unicodedata
import zipfile
from pathlib import PurePath

from docx import Document
from pypdf import PdfReader

from app.exceptions import AppError

MIME_TYPES = {
    ".txt": "text/plain",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


class ExtractionService:
    def __init__(self, max_chars: int):
        self.max_chars = max_chars

    def extract(self, content: bytes, filename: str, content_type: str) -> str:
        extension = PurePath(filename).suffix.lower()
        if extension not in MIME_TYPES or content_type.split(";")[0].lower() != MIME_TYPES[extension]:
            raise AppError("UNSUPPORTED_MEDIA", "Unsupported document format.", 415)
        if not content:
            raise AppError("EMPTY_DOCUMENT", "Document is empty.")
        try:
            if extension == ".txt":
                text = content.decode("utf-8-sig")
                if "\x00" in text:
                    raise ValueError("Binary text")
            elif extension == ".pdf":
                if not content.startswith(b"%PDF-"):
                    raise ValueError("Invalid PDF signature")
                reader = PdfReader(io.BytesIO(content), strict=True)
                if reader.is_encrypted:
                    raise ValueError("Encrypted PDF")
                parts = []
                total = 0
                for page in reader.pages:
                    part = page.extract_text() or ""
                    total += len(part)
                    if total > self.max_chars:
                        raise AppError("EXTRACTION_LIMIT", "Document text exceeds processing limit.", 413)
                    parts.append(part)
                text = "\n".join(parts)
            else:
                with zipfile.ZipFile(io.BytesIO(content)) as archive:
                    if sum(info.file_size for info in archive.infolist()) > max(
                        1_000_000, self.max_chars * 10
                    ):
                        raise AppError("EXTRACTION_LIMIT", "Document exceeds processing limit.", 413)
                    if "word/document.xml" not in archive.namelist():
                        raise ValueError("Invalid DOCX")
                document = Document(io.BytesIO(content))
                parts = [paragraph.text for paragraph in document.paragraphs]
                parts.extend(
                    cell.text for table in document.tables for row in table.rows for cell in row.cells
                )
                text = "\n".join(parts)
        except AppError:
            raise
        except Exception as exc:
            raise AppError("MALFORMED_DOCUMENT", "Document could not be parsed.") from exc
        text = unicodedata.normalize("NFKC", text).strip()
        if len(text) > self.max_chars:
            raise AppError("EXTRACTION_LIMIT", "Document text exceeds processing limit.", 413)
        if not text:
            raise AppError("EMPTY_DOCUMENT", "No readable text was found; OCR is not supported.")
        return text
